
import os
from datetime import datetime, timedelta

import openpyxl


# ============================================================
# CONFIGURATION
# ============================================================

SOURCE_FILE = r"C:\\Users\\trish\\Downloads\\iMoni_Reports_Summary_3.xlsx"
TEMPLATE_FILE = r"D:\\ATS\\server\\ATS\\template\\srms_allpassed_template.xlsx"

OUTPUT_FOLDER = r"D:\\TechnoTrendz\\ATS\\Generated_Reports"


# ============================================================
# DATE/TIME CONFIGURATION
# ============================================================
# Start generating report timestamps from:
START_DATETIME = datetime(2026, 9, 7, 9, 0)

# Working hours
WORK_START_HOUR = 9
WORK_START_MINUTE = 0

WORK_END_HOUR = 19
WORK_END_MINUTE = 0

# Gap between reports
TIME_GAP_MINUTES = 20


# ============================================================
# COLUMN MAPPING
# ============================================================

# Summary Excel columns:
#
# A = S.No.
# B = Report No
# C = Assembly Sr. No.
# D = DeviceIP
# E = CPU Sr. No.
# F = Base PCB Sr. No.
# G = Camera Sr. No.
# H = PSU Sr. No.

SUMMARY_COLUMNS = {
    "report_no": 2,
    "assembly_sr_no": 3,
    "device_ip": 4,
    "cpu_sr_no": 5,
    "base_pcb_sr_no": 6,
    "camera_sr_no": 7,
    "psu_sr_no": 8,
}


# ============================================================
# CLEAN VALUE
# ============================================================

def clean_value(value):
    """
    Convert empty Excel cells to blank strings.
    """
    if value is None:
        return ""

    return value


# ============================================================
# GET NEXT WORKING DATETIME
# ============================================================

def get_next_report_datetime(current_datetime):
    """
    Return the next valid report datetime.

    Rules:
    - Reports start at 09:00 AM.
    - Reports can be generated until 07:00 PM.
    - Gap between reports = 20 minutes.
    - Sunday is skipped.
    """

    next_datetime = current_datetime + timedelta(
        minutes=TIME_GAP_MINUTES
    )

    while True:

        # ----------------------------------------------------
        # Skip Sunday
        #
        # Monday = 0
        # Tuesday = 1
        # ...
        # Saturday = 5
        # Sunday = 6
        # ----------------------------------------------------

        if next_datetime.weekday() == 6:

            # Move to Monday
            next_datetime = (
                next_datetime.replace(
                    hour=WORK_START_HOUR,
                    minute=WORK_START_MINUTE,
                    second=0,
                    microsecond=0
                )
                + timedelta(days=1)
            )

            continue

        # ----------------------------------------------------
        # If time has gone beyond 7 PM
        # move to next working day at 9 AM
        # ----------------------------------------------------

        if (
            next_datetime.hour > WORK_END_HOUR
            or (
                next_datetime.hour == WORK_END_HOUR
                and next_datetime.minute > WORK_END_MINUTE
            )
        ):

            next_datetime = next_datetime + timedelta(days=1)

            next_datetime = next_datetime.replace(
                hour=WORK_START_HOUR,
                minute=WORK_START_MINUTE,
                second=0,
                microsecond=0
            )

            continue

        # ----------------------------------------------------
        # If somehow time is before working hours
        # ----------------------------------------------------

        if (
            next_datetime.hour < WORK_START_HOUR
            or (
                next_datetime.hour == WORK_START_HOUR
                and next_datetime.minute < WORK_START_MINUTE
            )
        ):

            next_datetime = next_datetime.replace(
                hour=WORK_START_HOUR,
                minute=WORK_START_MINUTE,
                second=0,
                microsecond=0
            )

            continue

        break

    return next_datetime


# ============================================================
# COPY TEMPLATE
# ============================================================

def copy_template():
    """
    Load a fresh copy of the template.
    """
    return openpyxl.load_workbook(TEMPLATE_FILE)


# ============================================================
# FILL REPORT
# ============================================================

def fill_report(template_wb, data, report_datetime):
    """
    Fill one report using the template.
    """

    ws = template_wb.active

    # --------------------------------------------------------
    # Header information
    # --------------------------------------------------------

    ws["B2"] = clean_value(data["report_no"])

    ws["B3"] = clean_value(
        data["assembly_sr_no"]
    )

    # --------------------------------------------------------
    # Date/time
    # --------------------------------------------------------

    ws["B4"] = report_datetime

    ws["B4"].number_format = "dd-mm-yyyy hh:mm:ss"

    # --------------------------------------------------------
    # Test level
    # --------------------------------------------------------

    ws["B5"] = "iMoni Assembly"

    # --------------------------------------------------------
    # Device IP
    # --------------------------------------------------------

    ws["B6"] = clean_value(
        data["device_ip"]
    )

    # --------------------------------------------------------
    # Hardware serial numbers
    # --------------------------------------------------------

    ws["C7"] = clean_value(
        data["cpu_sr_no"]
    )

    ws["D7"] = clean_value(
        data["base_pcb_sr_no"]
    )

    ws["E7"] = clean_value(
        data["camera_sr_no"]
    )

    ws["F7"] = clean_value(
        data["psu_sr_no"]
    )

    # --------------------------------------------------------
    # Total tests
    # --------------------------------------------------------

    ws["B7"] = 16

    # --------------------------------------------------------
    # Make sure all 16 tests are PASSED
    #
    # Test rows = 10 to 25
    # Column C = Status
    # --------------------------------------------------------

    for row in range(10, 26):
        ws.cell(
            row=row,
            column=3
        ).value = "PASSED"

    return template_wb


# ============================================================
# SAFE FILENAME
# ============================================================

def safe_filename(filename):
    """
    Remove characters that are invalid in Windows filenames.
    """

    invalid_chars = '<>:"/\\|?*'

    for char in invalid_chars:
        filename = filename.replace(
            char,
            "_"
        )

    return filename.strip()


# ============================================================
# GENERATE REPORTS
# ============================================================

def generate_reports():

    # --------------------------------------------------------
    # Check files
    # --------------------------------------------------------

    if not os.path.exists(SOURCE_FILE):
        raise FileNotFoundError(
            f"Source file not found: {SOURCE_FILE}"
        )

    if not os.path.exists(TEMPLATE_FILE):
        raise FileNotFoundError(
            f"Template file not found: {TEMPLATE_FILE}"
        )

    # --------------------------------------------------------
    # Create output folder
    # --------------------------------------------------------

    os.makedirs(
        OUTPUT_FOLDER,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Load summary Excel
    # --------------------------------------------------------

    source_wb = openpyxl.load_workbook(
        SOURCE_FILE,
        data_only=False
    )

    source_ws = source_wb.active

    generated = 0
    skipped = 0

    # --------------------------------------------------------
    # Current report datetime
    #
    # First report will be exactly:
    # 07-09-2026 09:00:00
    # --------------------------------------------------------

    current_report_datetime = START_DATETIME

    # --------------------------------------------------------
    # Process every row
    # --------------------------------------------------------

    for row in range(
        2,
        source_ws.max_row + 1
    ):

        report_no = source_ws.cell(
            row=row,
            column=SUMMARY_COLUMNS["report_no"]
        ).value

        assembly_sr_no = source_ws.cell(
            row=row,
            column=SUMMARY_COLUMNS["assembly_sr_no"]
        ).value

        # ----------------------------------------------------
        # Skip completely empty rows
        # ----------------------------------------------------

        if not report_no and not assembly_sr_no:
            skipped += 1
            continue

        # ----------------------------------------------------
        # Read data
        # ----------------------------------------------------

        data = {
            "report_no": source_ws.cell(
                row=row,
                column=SUMMARY_COLUMNS["report_no"]
            ).value,

            "assembly_sr_no": source_ws.cell(
                row=row,
                column=SUMMARY_COLUMNS["assembly_sr_no"]
            ).value,

            "device_ip": source_ws.cell(
                row=row,
                column=SUMMARY_COLUMNS["device_ip"]
            ).value,

            "cpu_sr_no": source_ws.cell(
                row=row,
                column=SUMMARY_COLUMNS["cpu_sr_no"]
            ).value,

            "base_pcb_sr_no": source_ws.cell(
                row=row,
                column=SUMMARY_COLUMNS["base_pcb_sr_no"]
            ).value,

            "camera_sr_no": source_ws.cell(
                row=row,
                column=SUMMARY_COLUMNS["camera_sr_no"]
            ).value,

            "psu_sr_no": source_ws.cell(
                row=row,
                column=SUMMARY_COLUMNS["psu_sr_no"]
            ).value,
        }

        # ----------------------------------------------------
        # Load fresh template
        # ----------------------------------------------------

        report_wb = copy_template()

        # ----------------------------------------------------
        # Fill report
        # ----------------------------------------------------

        report_wb = fill_report(
            report_wb,
            data,
            current_report_datetime
        )

        # ----------------------------------------------------
        # Filename
        # ----------------------------------------------------

        filename = (
            f"{data['report_no']}_"
            f"{data['assembly_sr_no']}.xlsx"
        )

        filename = safe_filename(filename)

        output_path = os.path.join(
            OUTPUT_FOLDER,
            filename
        )

        # ----------------------------------------------------
        # Save
        # ----------------------------------------------------

        report_wb.save(output_path)

        generated += 1

        print(
            f"[{generated}] "
            f"{current_report_datetime.strftime('%d-%m-%Y %H:%M:%S')} "
            f"-> Generated: {filename}"
        )

        # ----------------------------------------------------
        # Calculate next report datetime
        # ----------------------------------------------------

        current_report_datetime = (
            get_next_report_datetime(
                current_report_datetime
            )
        )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    print()
    print("=" * 60)
    print("REPORT GENERATION COMPLETE")
    print("=" * 60)

    print(
        f"Reports generated : {generated}"
    )

    print(
        f"Rows skipped      : {skipped}"
    )

    print(
        f"Output folder     : "
        f"{os.path.abspath(OUTPUT_FOLDER)}"
    )

    print("=" * 60)


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    generate_reports()
